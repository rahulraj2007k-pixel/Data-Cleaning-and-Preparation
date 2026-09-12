from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


# ---------------------------------------------------------
# PROJECT PATHS
# ---------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
REPORT_DIR = BASE_DIR / "report"
SCREENSHOTS_DIR = BASE_DIR / "screenshots"

REPORT_DIR.mkdir(exist_ok=True)

OUTPUT_FILE = REPORT_DIR / "Data_Cleaning_Report.docx"


# ---------------------------------------------------------
# DOCUMENT SETUP
# ---------------------------------------------------------
doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)


# ---------------------------------------------------------
# STYLES
# ---------------------------------------------------------
styles = doc.styles

styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)

styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True

styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.size = Pt(17)
styles["Heading 1"].font.bold = True

styles["Heading 2"].font.name = "Arial"
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True


# ---------------------------------------------------------
# HELPER FUNCTIONS
# ---------------------------------------------------------
def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "

    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_number(text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_screenshot(filename, caption, width=6.2):
    path = SCREENSHOTS_DIR / filename

    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = p.add_run()
        run.add_picture(str(path), width=Inches(width))

        caption_p = doc.add_paragraph()
        caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        r = caption_p.add_run(caption)
        r.bold = True
        r.font.size = Pt(9)


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells

    for i, header in enumerate(headers):
        header_cells[i].text = str(header)
        header_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for run in header_cells[i].paragraphs[0].runs:
            run.bold = True

    for row in rows:
        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    return table


# Footer
footer = section.footer
add_page_number(footer.paragraphs[0])


# ---------------------------------------------------------
# TITLE PAGE
# ---------------------------------------------------------
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.space_after = Pt(40)

r = p.add_run("DATA CLEANING AND PREPARATION")
r.bold = True
r.font.size = Pt(26)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run("Project Report")
r.bold = True
r.font.size = Pt(18)

doc.add_paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run(
    "A Python-based data cleaning, validation,\n"
    "preparation and analysis project"
)
r.font.size = Pt(14)

doc.add_paragraph("")
doc.add_paragraph("")
doc.add_paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run("Technology Stack")
r.bold = True
r.font.size = Pt(13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Python  •  Pandas  •  NumPy  •  Matplotlib").font.size = Pt(11)

doc.add_paragraph("")
doc.add_paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run("GitHub Repository")
r.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(
    "https://github.com/rahulraj2007k-pixel/"
    "Data-Cleaning-and-Preparation"
)

doc.add_page_break()


# ---------------------------------------------------------
# 1. INTRODUCTION
# ---------------------------------------------------------
doc.add_heading("1. Introduction", level=1)

doc.add_paragraph(
    "Data cleaning and preparation is an important stage of the data "
    "analysis process. Real-world datasets often contain missing values, "
    "invalid records, inconsistent formatting, duplicate records and "
    "incorrect data types. If these problems are not addressed, they can "
    "produce unreliable analysis results."
)

doc.add_paragraph(
    "This project demonstrates a complete data cleaning and preparation "
    "workflow using Python. A synthetic customer sales dataset was created "
    "with realistic data-quality problems. The dataset was then cleaned, "
    "validated and prepared for analysis."
)


# ---------------------------------------------------------
# 2. OBJECTIVES
# ---------------------------------------------------------
doc.add_heading("2. Project Objectives", level=1)

objectives = [
    "Load and inspect a raw CSV dataset.",
    "Identify missing, invalid and inconsistent data.",
    "Standardize text fields and column names.",
    "Validate email addresses.",
    "Convert multiple date formats into a consistent format.",
    "Handle invalid negative quantity and price values.",
    "Detect and remove business-key duplicate records.",
    "Fill missing values using appropriate strategies.",
    "Create a calculated total_amount field.",
    "Validate the cleaned dataset.",
    "Perform basic sales analysis and generate visualizations."
]

for item in objectives:
    add_bullet(item)


# ---------------------------------------------------------
# 3. DATASET DESCRIPTION
# ---------------------------------------------------------
doc.add_heading("3. Dataset Description", level=1)

doc.add_paragraph(
    "The project uses a synthetic customer sales dataset created specifically "
    "for demonstrating data cleaning techniques. It does not contain private "
    "customer information or confidential records."
)

add_table(
    ["Property", "Value"],
    [
        ["Original rows", "30"],
        ["Original columns", "8"],
        ["Final rows", "29"],
        ["Final columns", "9"],
        ["Input format", "CSV"],
        ["Output format", "CSV"],
    ],
)

doc.add_paragraph("")

doc.add_heading("Dataset Columns", level=2)

add_table(
    ["Column", "Description"],
    [
        ["order_id", "Unique order identifier"],
        ["customer_name", "Customer name"],
        ["email", "Customer email address"],
        ["city", "Customer city"],
        ["order_date", "Date of order"],
        ["product", "Purchased product"],
        ["quantity", "Quantity purchased"],
        ["unit_price", "Price per unit"],
        ["total_amount", "Calculated total sales amount"],
    ],
)


# ---------------------------------------------------------
# 4. DATA QUALITY ISSUES
# ---------------------------------------------------------
doc.add_heading("4. Identified Data Quality Issues", level=1)

issues = [
    ["Missing email", "2", "Handled using a placeholder"],
    ["Missing order date", "1", "Filled using median date"],
    ["Missing quantity", "1", "Filled using median"],
    ["Missing unit price", "1", "Filled using median"],
    ["Invalid email format", "1", "Converted to missing and handled"],
    ["Negative quantity", "1", "Converted to missing"],
    ["Negative unit price", "1", "Converted to missing"],
    ["Inconsistent text formatting", "Multiple", "Standardized"],
    ["Multiple date formats", "Multiple", "Converted to standard date"],
    ["Business-key duplicate", "1", "Removed"],
]

add_table(
    ["Issue", "Count / Scope", "Action Taken"],
    issues
)


# ---------------------------------------------------------
# 5. CLEANING METHODOLOGY
# ---------------------------------------------------------
doc.add_heading("5. Data Cleaning Methodology", level=1)

steps = [
    "Load the raw CSV dataset using Pandas.",
    "Inspect dataset shape, columns, data types and missing values.",
    "Standardize column names.",
    "Remove leading and trailing spaces from text values.",
    "Standardize city and product capitalization.",
    "Convert email addresses to lowercase.",
    "Validate email addresses using a regular expression.",
    "Parse multiple date formats using Pandas.",
    "Convert quantity and unit price into numeric data types.",
    "Convert negative quantities and prices into missing values.",
    "Detect business-key duplicates.",
    "Remove duplicate business records while preserving the first occurrence.",
    "Fill missing quantity and price values using the median.",
    "Fill missing emails using a safe placeholder.",
    "Fill missing order dates using the median date.",
    "Create the total_amount calculated field.",
    "Perform final validation and save the cleaned dataset."
]

for i, item in enumerate(steps, start=1):
    add_number(item)


# ---------------------------------------------------------
# 6. IMPLEMENTATION
# ---------------------------------------------------------
doc.add_heading("6. Implementation", level=1)

doc.add_paragraph(
    "The implementation is divided into two Python scripts. "
    "The first script performs cleaning and preparation, while the second "
    "script performs analysis and visualization."
)

doc.add_heading("6.1 Data Cleaning Script", level=2)

doc.add_paragraph(
    "File: src/data_cleaning.py"
)

add_bullet("Loads data/raw_dataset.csv.")
add_bullet("Performs validation and transformation.")
add_bullet("Generates cleaning_log.txt.")
add_bullet("Creates data/cleaned_dataset.csv.")
add_bullet("Reports final missing values and duplicate counts.")


doc.add_heading("6.2 Data Analysis Script", level=2)

doc.add_paragraph(
    "File: src/data_analysis.py"
)

add_bullet("Compares raw and cleaned datasets.")
add_bullet("Compares missing values before and after cleaning.")
add_bullet("Calculates descriptive statistics.")
add_bullet("Calculates product-wise sales.")
add_bullet("Calculates city-wise sales.")
add_bullet("Generates charts using Matplotlib.")
add_bullet("Creates analysis_summary.txt.")


# ---------------------------------------------------------
# 7. RESULTS
# ---------------------------------------------------------
doc.add_heading("7. Cleaning Results", level=1)

add_table(
    ["Metric", "Before Cleaning", "After Cleaning"],
    [
        ["Rows", "30", "29"],
        ["Columns", "8", "9"],
        ["Missing values", "Present", "0"],
        ["Exact duplicate rows", "0", "0"],
        ["Business-key duplicates", "1", "0"],
        ["Invalid emails", "1", "0"],
        ["Invalid negative values", "2", "0"],
    ],
)

doc.add_paragraph("")

doc.add_paragraph(
    "The final dataset contains 29 rows and 9 columns. All missing values "
    "were handled, invalid values were corrected, and duplicate business "
    "records were removed."
)


# ---------------------------------------------------------
# 8. ANALYSIS
# ---------------------------------------------------------
doc.add_heading("8. Data Analysis and Visualization", level=1)

doc.add_paragraph(
    "After cleaning, the prepared dataset was analyzed to demonstrate that "
    "the resulting data can be used reliably for basic business analysis."
)

doc.add_heading("8.1 Product-wise Sales", level=2)

add_table(
    ["Product", "Total Sales"],
    [
        ["Laptop", "439,000"],
        ["Tablet", "256,500"],
        ["Mobile Phone", "211,500"],
        ["Headphones", "20,800"],
    ],
)

doc.add_paragraph("")

add_screenshot(
    "product_sales.png",
    "Figure 1: Product-wise sales visualization"
)


doc.add_heading("8.2 City-wise Sales", level=2)

add_table(
    ["City", "Total Sales"],
    [
        ["Delhi", "397,800"],
        ["Bangalore", "229,000"],
        ["Mumbai", "195,100"],
        ["Kolkata", "105,900"],
    ],
)

doc.add_paragraph("")

add_screenshot(
    "city_sales.png",
    "Figure 2: City-wise sales visualization"
)


# ---------------------------------------------------------
# 9. BEFORE/AFTER VISUALIZATION
# ---------------------------------------------------------
doc.add_heading("9. Data Quality Comparison", level=1)

doc.add_paragraph(
    "The following visualization compares missing values in the raw and "
    "cleaned datasets."
)

add_screenshot(
    "missing_values_comparison.png",
    "Figure 3: Missing values before and after cleaning"
)


# ---------------------------------------------------------
# 10. CLEANING OUTPUT
# ---------------------------------------------------------
doc.add_heading("10. Cleaning Execution Output", level=1)

doc.add_paragraph(
    "The cleaning process was executed successfully from the terminal. "
    "The following screenshot records the execution and validation output."
)

add_screenshot(
    "cleaning_output.png",
    "Figure 4: Data cleaning execution output"
)


# ---------------------------------------------------------
# 11. VALIDATION
# ---------------------------------------------------------
doc.add_heading("11. Validation", level=1)

validation = [
    "Final missing value count: 0.",
    "Final exact duplicate count: 0.",
    "Final dataset shape: 29 rows × 9 columns.",
    "Quantity values are stored as integers.",
    "Unit prices are stored as floating-point numbers.",
    "The total_amount field is calculated from quantity × unit_price.",
    "Dates are standardized to YYYY-MM-DD format.",
    "Invalid email records are handled safely."
]

for item in validation:
    add_bullet(item)


# ---------------------------------------------------------
# 12. TOOLS AND LIBRARIES
# ---------------------------------------------------------
doc.add_heading("12. Tools and Libraries", level=1)

add_table(
    ["Tool / Library", "Purpose"],
    [
        ["Python", "Programming language"],
        ["Pandas", "Data loading, cleaning and analysis"],
        ["NumPy", "Numerical operations"],
        ["Matplotlib", "Data visualization"],
        ["python-docx", "Report generation"],
        ["Git", "Version control"],
        ["GitHub", "Source code hosting"],
        ["VS Code / Terminal", "Development environment"],
    ],
)


# ---------------------------------------------------------
# 13. PROJECT STRUCTURE
# ---------------------------------------------------------
doc.add_heading("13. Project Structure", level=1)

structure = """Data-Cleaning-and-Preparation/
│
├── data/
│   ├── raw_dataset.csv
│   └── cleaned_dataset.csv
│
├── report/
│   └── Data_Cleaning_Report.docx
│
├── screenshots/
│   ├── cleaning_output.png
│   ├── missing_values_comparison.png
│   ├── product_sales.png
│   └── city_sales.png
│
├── src/
│   ├── data_cleaning.py
│   └── data_analysis.py
│
├── analysis_summary.txt
├── cleaning_log.txt
├── README.md
├── requirements.txt
└── .gitignore"""

p = doc.add_paragraph()
r = p.add_run(structure)
r.font.name = "Consolas"
r.font.size = Pt(9)


# ---------------------------------------------------------
# 14. SETUP AND USAGE
# ---------------------------------------------------------
doc.add_heading("14. Setup and Usage", level=1)

doc.add_paragraph(
    "Install the required Python libraries using:"
)

p = doc.add_paragraph()
r = p.add_run("pip install -r requirements.txt")
r.font.name = "Consolas"
r.bold = True

doc.add_paragraph(
    "Run the data cleaning process:"
)

p = doc.add_paragraph()
r = p.add_run("python src/data_cleaning.py")
r.font.name = "Consolas"
r.bold = True

doc.add_paragraph(
    "Run the analysis and visualization process:"
)

p = doc.add_paragraph()
r = p.add_run("python src/data_analysis.py")
r.font.name = "Consolas"
r.bold = True


# ---------------------------------------------------------
# 15. RESOURCES
# ---------------------------------------------------------
doc.add_heading("15. External Libraries, APIs, Datasets and Resources", level=1)

add_bullet(
    "Dataset: Synthetic customer sales dataset created specifically "
    "for this project."
)

add_bullet(
    "Pandas: Used for data manipulation and cleaning."
)

add_bullet(
    "NumPy: Used for numerical data processing."
)

add_bullet(
    "Matplotlib: Used for charts and visualizations."
)

add_bullet(
    "python-docx: Used to generate this project report."
)

add_bullet(
    "No external API is required for this project."
)

add_bullet(
    "No private credentials, passwords or API keys are included."
)


# ---------------------------------------------------------
# 16. SECURITY AND PRIVACY
# ---------------------------------------------------------
doc.add_heading("16. Security and Privacy", level=1)

doc.add_paragraph(
    "The dataset used in this project is synthetic and does not represent "
    "real customer records. The repository does not contain passwords, "
    "API keys, private credentials or other sensitive authentication data."
)


# ---------------------------------------------------------
# 17. CONCLUSION
# ---------------------------------------------------------
doc.add_heading("17. Conclusion", level=1)

doc.add_paragraph(
    "This project demonstrates a complete data cleaning and preparation "
    "workflow using Python. The raw dataset contained realistic data-quality "
    "issues such as missing values, invalid emails, inconsistent formatting, "
    "multiple date formats, negative numerical values and a business-key "
    "duplicate."
)

doc.add_paragraph(
    "After applying systematic cleaning and validation techniques, the "
    "dataset was transformed into a consistent 29-row, 9-column dataset "
    "with zero remaining missing values and zero exact duplicate rows. "
    "The cleaned data was then used to produce product-wise and city-wise "
    "sales analysis."
)

doc.add_paragraph(
    "The project provides a reproducible foundation for further analytics "
    "and demonstrates practical data preparation skills using Python."
)


# ---------------------------------------------------------
# 18. GITHUB
# ---------------------------------------------------------
doc.add_heading("18. GitHub Repository", level=1)

p = doc.add_paragraph()
p.add_run(
    "https://github.com/rahulraj2007k-pixel/"
    "Data-Cleaning-and-Preparation"
)


# ---------------------------------------------------------
# AUTHOR
# ---------------------------------------------------------
doc.add_heading("19. Author", level=1)

doc.add_paragraph(
    "Student Project — Data Cleaning and Preparation"
)

doc.add_paragraph(
    "The project source code, datasets, analysis files, screenshots and "
    "documentation are maintained in the GitHub repository listed above."
)


# ---------------------------------------------------------
# SAVE
# ---------------------------------------------------------
doc.save(OUTPUT_FILE)

print()
print("=" * 60)
print("REPORT GENERATED SUCCESSFULLY")
print("=" * 60)
print(f"File: {OUTPUT_FILE}")
print("=" * 60)